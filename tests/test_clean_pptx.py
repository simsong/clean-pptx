import io
import shutil
import tempfile
import unittest
from contextlib import redirect_stderr
from pathlib import Path
from uuid import uuid4

import clean_pptx


class TestCleanPptx(unittest.TestCase):
    def test_derive_output_paths(self):
        pptx_path = Path("demo deck.pptx")
        cleaned_pptx, cleaned_pdf = clean_pptx.derive_output_paths(pptx_path)
        self.assertEqual(cleaned_pptx, Path("demo deck-cleaned.pptx"))
        self.assertEqual(cleaned_pdf, Path("demo deck-cleaned.pdf"))

    def test_parse_args_accepts_skip_pdf(self):
        args = clean_pptx.parse_args(["slides.pptx", "--skip-pdf"])
        self.assertEqual(args.pptx_file, Path("slides.pptx"))
        self.assertTrue(args.skip_pdf)

    def test_main_rejects_non_pptx_extension(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            input_file = Path(tmpdir) / "file.txt"
            input_file.write_text("not a pptx", encoding="utf-8")
            cleaned_pptx, cleaned_pdf = clean_pptx.derive_output_paths(input_file)

            stderr = io.StringIO()
            with redirect_stderr(stderr):
                rc = clean_pptx.main([str(input_file)])

        self.assertEqual(rc, 2)
        self.assertIn("Error: input file must have a .pptx extension.", stderr.getvalue())
        self.assertFalse(cleaned_pptx.exists())
        self.assertFalse(cleaned_pdf.exists())

    def test_main_rejects_missing_pptx_file(self):
        missing_name = f"missing-{uuid4()}.pptx"
        stderr = io.StringIO()
        with redirect_stderr(stderr):
            rc = clean_pptx.main([missing_name])
        self.assertEqual(rc, 2)
        self.assertIn("Error: file not found:", stderr.getvalue())

    def test_main_cleans_sample_pptx(self):
        fixture = Path(__file__).parent / "fixtures" / "sample.pptx"
        if not fixture.exists():
            self.skipTest("sample fixture is not available")

        try:
            from pptx import Presentation
        except ImportError:
            self.skipTest("python-pptx not installed")

        with tempfile.TemporaryDirectory() as tmpdir:
            tmp_input = Path(tmpdir) / "sample.pptx"
            shutil.copyfile(fixture, tmp_input)
            rc = clean_pptx.main([str(tmp_input), "--skip-pdf"])
            self.assertEqual(rc, 0)

            cleaned_pptx, _ = clean_pptx.derive_output_paths(tmp_input)
            self.assertTrue(cleaned_pptx.exists())
            self.assertGreater(cleaned_pptx.stat().st_size, 0)

            original_slides = len(Presentation(str(tmp_input)).slides)
            cleaned_slides = len(Presentation(str(cleaned_pptx)).slides)
            self.assertEqual(cleaned_slides, original_slides)


if __name__ == "__main__":
    unittest.main()
