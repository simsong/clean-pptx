#!/usr/bin/env python3
"""Clean PPTX decks and render cleaned PDF output via Microsoft PowerPoint on macOS."""

from __future__ import annotations

import argparse
import io
import subprocess
import sys
import warnings
from pathlib import Path
from typing import Iterable

STANDARD_FONT_MAP = {
    "Calibri": "Arial",
    "Cambria": "Times New Roman",
    "Aptos": "Arial",
    "Arial": "Arial",
    "Times New Roman": "Times New Roman",
    "Courier New": "Courier New",
}
DEFAULT_STANDARD_FONT = "Arial"


def derive_output_paths(input_path: Path) -> tuple[Path, Path]:
    stem = input_path.stem
    cleaned_pptx = input_path.with_name(f"{stem}-cleaned.pptx")
    cleaned_pdf = input_path.with_name(f"{stem}-cleaned.pdf")
    return cleaned_pptx, cleaned_pdf


def _standard_font_name(name: str | None) -> str:
    if not name:
        return DEFAULT_STANDARD_FONT
    return STANDARD_FONT_MAP.get(name, DEFAULT_STANDARD_FONT)


def _copy_text_frame(source, dest) -> None:
    dest.clear()
    for idx, src_paragraph in enumerate(source.paragraphs):
        paragraph = dest.paragraphs[0] if idx == 0 else dest.add_paragraph()
        paragraph.level = src_paragraph.level
        paragraph.alignment = src_paragraph.alignment
        for src_run in src_paragraph.runs:
            run = paragraph.add_run()
            run.text = src_run.text
            run.font.name = _standard_font_name(src_run.font.name)
            run.font.bold = src_run.font.bold
            run.font.italic = src_run.font.italic
            run.font.underline = src_run.font.underline
            if src_run.font.size:
                run.font.size = src_run.font.size
            if src_run.font.color and src_run.font.color.type and src_run.font.color.rgb:
                run.font.color.rgb = src_run.font.color.rgb


def _reencode_to_jpeg(image_blob: bytes, quality: int = 90) -> io.BytesIO:
    """Convert an image blob to JPEG and return a BytesIO object positioned at start."""
    from PIL import Image

    with Image.open(io.BytesIO(image_blob)) as img:
        rgb = img.convert("RGB")
        output = io.BytesIO()
        rgb.save(output, format="JPEG", quality=quality, optimize=True)
        output.seek(0)
        return output


def _iter_slide_shapes(slide, group_shape_type) -> Iterable:
    for shape in slide.shapes:
        if shape.shape_type == group_shape_type:
            for subshape in shape.shapes:
                yield subshape
        else:
            yield shape


def create_cleaned_pptx(input_path: Path, output_path: Path) -> None:
    """Rebuild a source PPTX into a cleaned PPTX with standard fonts and JPEG images."""
    try:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE
        from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
    except ImportError as exc:
        raise RuntimeError(
            "python-pptx is required. Install with: pip install python-pptx"
        ) from exc

    source = Presentation(str(input_path))
    cleaned = Presentation()

    if cleaned.slides:
        r_id = cleaned.slides._sldIdLst[0].rId
        cleaned.part.drop_rel(r_id)
        del cleaned.slides._sldIdLst[0]

    blank_layout = cleaned.slide_layouts[6]

    for src_slide in source.slides:
        dest_slide = cleaned.slides.add_slide(blank_layout)

        for shape in _iter_slide_shapes(src_slide, MSO_SHAPE_TYPE.GROUP):
            left = shape.left
            top = shape.top
            width = shape.width
            height = shape.height

            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                jpeg_data = _reencode_to_jpeg(shape.image.blob, quality=90)
                dest_slide.shapes.add_picture(jpeg_data, left, top, width=width, height=height)
                continue

            if (
                shape.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE
                and getattr(shape, "auto_shape_type", None) is not None
            ):
                try:
                    auto_shape = MSO_AUTO_SHAPE_TYPE(shape.auto_shape_type)
                except ValueError:
                    auto_shape = MSO_AUTO_SHAPE_TYPE.RECTANGLE
                new_shape = dest_slide.shapes.add_shape(auto_shape, left, top, width, height)
                if getattr(shape, "has_text_frame", False):
                    _copy_text_frame(shape.text_frame, new_shape.text_frame)
                continue

            if getattr(shape, "has_text_frame", False):
                textbox = dest_slide.shapes.add_textbox(left, top, width, height)
                _copy_text_frame(shape.text_frame, textbox.text_frame)
                continue

            # Unsupported shape types are replaced by empty rectangles to preserve position.
            warnings.warn(
                f"Unsupported shape type {shape.shape_type!r} replaced with rectangle.",
                RuntimeWarning,
            )
            dest_slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, left, top, width, height)

    cleaned.save(str(output_path))


def render_pdf_with_powerpoint(cleaned_pptx_path: Path, output_pdf_path: Path) -> None:
    """Render a PPTX to PDF by driving Microsoft PowerPoint via AppleScript on macOS."""
    if sys.platform != "darwin":
        raise RuntimeError("PDF rendering requires macOS with Microsoft PowerPoint installed.")

    applescript = """
    on run argv
        set pptxPath to item 1 of argv
        set pdfPath to item 2 of argv
        tell application "Microsoft PowerPoint"
            activate
            set thePresentation to open POSIX file pptxPath
            save as thePresentation file name POSIX file pdfPath file format save as PDF
            close thePresentation saving no
        end tell
    end run
    """

    try:
        subprocess.run(
            ["osascript", "-e", applescript, str(cleaned_pptx_path.resolve()), str(output_pdf_path.resolve())],
            check=True,
            capture_output=True,
            text=True,
        )
    except subprocess.CalledProcessError as exc:
        details = (exc.stderr or exc.stdout or "").strip()
        raise RuntimeError(
            f"PowerPoint PDF rendering failed. Ensure Microsoft PowerPoint is installed and scriptable. {details}"
        ) from exc


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Create <original>-cleaned.pptx and render <original>-cleaned.pdf "
            "using Microsoft PowerPoint on macOS."
        )
    )
    parser.add_argument("pptx_file", type=Path, help="Path to source .pptx file")
    parser.add_argument(
        "--skip-pdf",
        action="store_true",
        help="Only generate cleaned PPTX (skip PowerPoint PDF rendering step)",
    )
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    input_path = args.pptx_file

    if input_path.suffix.lower() != ".pptx":
        print("Error: input file must have a .pptx extension.", file=sys.stderr)
        return 2

    if not input_path.exists():
        print(f"Error: file not found: {input_path}", file=sys.stderr)
        return 2

    cleaned_pptx, cleaned_pdf = derive_output_paths(input_path)

    try:
        create_cleaned_pptx(input_path, cleaned_pptx)
        if not args.skip_pdf:
            render_pdf_with_powerpoint(cleaned_pptx, cleaned_pdf)
    except (RuntimeError, OSError, ValueError) as exc:
        print(f"Error ({type(exc).__name__}): {exc}", file=sys.stderr)
        return 1

    print(f"Created cleaned deck: {cleaned_pptx}")
    if args.skip_pdf:
        print("Skipped PDF rendering (--skip-pdf).")
    else:
        print(f"Created rendered PDF: {cleaned_pdf}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
